"""Collect the fourteen cached fits and export one sheet per run plus report."""
import json
import platform
from pathlib import Path
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from pilot.common.run import dump
from pilot.b1_fourteen.run import OUT, WORK, MODELS, VERSION, GROUPS

PERIODS = ['۱۳۹۲–۱۳۹۶','۱۳۹۷–۱۳۹۸','۱۳۹۹–۱۴۰۳']


def mdtable(headers, rows):
    lines = ['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']
    for row in rows:
        lines.append('| '+' | '.join(f'{v:.6g}' if isinstance(v,(float,np.floating)) else str(v) for v in row)+' |')
    return '\n'.join(lines)+'\n'


def model_description(spec):
    if not spec['panel']: kind = 'مقطع تجمیعی؛ بدون میانگین موندلاک'
    elif spec['cohort']: kind = f"پنل سه‌سالهٔ مستقل گروه ورود {spec['cohort']}"
    else: kind = 'پنل تجمیعی هشت گروه ورود؛ موندلاک'
    return kind+('؛ ضرایب β، Γ، λ متفاوت در سه دورهٔ تاریخی' if spec['history'] else '؛ ضرایب مشترک')


def economic_comparisons(results):
    byname={r['specification']['name']:r for r in results}
    summary={}
    for label,a,b in [('Panel_B_vs_CS_B','Panel_B_common','CS_B_common'),
                      ('CS_All_vs_CS_B','CS_All_common','CS_B_common')]:
        ra=byname[a]['elasticities']['overall']; rb=byname[b]['elasticities']['overall']
        summary[label]={'own_price_difference':np.diag(np.array(ra['marshallian'])-np.array(rb['marshallian'])),
                        'expenditure_difference':np.array(ra['expenditure'])-np.array(rb['expenditure']),
                        'interpretation':'descriptive; model/sample/control differences are not an isolated causal CRE effect'}
    cohorts=[r for r in results if r['specification']['cohort'] is not None]
    weight=np.array([r['sample']['households'] for r in cohorts],float); weight/=weight.sum()
    own=np.array([np.diag(r['elasticities']['overall']['marshallian']) for r in cohorts])
    ex=np.array([r['elasticities']['overall']['expenditure'] for r in cohorts])
    summary['separate_cohorts']={'own_min':own.min(axis=0),'own_max':own.max(axis=0),
        'own_equal_cohort_mean':own.mean(axis=0),'own_household_count_weighted_mean':weight@own,
        'expenditure_min':ex.min(axis=0),'expenditure_max':ex.max(axis=0),
        'expenditure_equal_cohort_mean':ex.mean(axis=0),'expenditure_household_count_weighted_mean':weight@ex,
        'interpretation':'both parameters and cohort economic evaluation conditions differ'}
    for r in results:
        if not r['specification']['history']: continue
        xx=np.array([r['elasticities'][f'common_reference_period_{c}']['expenditure'] for c in range(3)])
        pp=np.array([np.diag(r['elasticities'][f'common_reference_period_{c}']['marshallian']) for c in range(3)])
        summary[r['specification']['name']]={'common_reference_own_by_period':pp,
             'common_reference_expenditure_by_period':xx,'own_range':np.ptp(pp,axis=0),
             'expenditure_range':np.ptp(xx,axis=0),
             'interpretation':'same global prices/expenditure/levels/means/CF; demand and participation regime responses differ'}
    return summary


def workbook(results):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    navy = '203864'; teal = 'DDEBF7'; green = 'E2F0D9'
    thin = Side(style='thin',color='B8C4D1')
    def merged(ws,row,text,color=None):
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=14)
        c=ws.cell(row,1,text); c.font=Font(name='Calibri',size=12,bold=bool(color),color='FFFFFF' if color else '000000')
        c.alignment=Alignment(horizontal='right',vertical='center',wrap_text=True)
        if color: c.fill=PatternFill('solid',fgColor=color)
        ws.row_dimensions[row].height=30 if len(text)>95 else 23
    def matrix(ws,row,ref,title):
        merged(ws,row,title,navy); row+=1
        for col,val in enumerate(['گروه / قیمت']+GROUPS+['کشش مخارج'],1):
            c=ws.cell(row,col,val); c.fill=PatternFill('solid',fgColor=teal)
            c.font=Font(name='Calibri',bold=True); c.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center')
        ws.row_dimensions[row].height=58
        for i in range(12):
            rr=row+i+1; ws.cell(rr,1,GROUPS[i])
            for j in range(12):
                value=ref['marshallian'][i][j]
                c=ws.cell(rr,j+2,float(value) if value is not None and np.isfinite(value) else 'نامعتبر')
                c.number_format='0.0000'; c.fill=PatternFill('solid',fgColor=green if i==j else 'FFFFFF')
            value=ref['expenditure'][i]
            ws.cell(rr,14,float(value) if value is not None and np.isfinite(value) else 'نامعتبر').number_format='0.0000'
            for cell in ws[rr]:
                cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
                cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
            ws.row_dimensions[rr].height=28
        return row+14
    for run in results:
        spec=run['specification']; ws=wb.create_sheet(spec['name'])
        ws.sheet_view.rightToLeft=True; ws.freeze_panes='B15'
        merged(ws,1,spec['name']+' — '+model_description(spec),navy)
        merged(ws,2,f"نمونه: {spec['sample']}؛ {run['sample']['rows']:,} مشاهده؛ {run['sample']['households']:,} خانوار؛ سال‌ها: "+'، '.join(map(str,run['sample']['years'])))
        merged(ws,3,'۱۲ گروه، ۱۰۶ قلم؛ نوشابه حذف شده؛ قیمت بازار مشترک با اهداکنندگان All؛ Young با وزن ثابت')
        penalty=[str(d['group']) for d in run['selection'] if d.get('estimation_method',{}).get('method')=='Jeffreys_prior_penalized_probit']
        merged(ws,4,'مدل: QUAIDS + S&Y دو مرحله‌ای + CF مخارج؛ مشارکت: پروبیت تجمیعی'+(' با میانگین‌های موندلاک' if spec['panel'] else '')+('؛ Jeffreys فقط گروه‌های '+','.join(penalty) if penalty else ''))
        years=[x for x in run['level_names'] if x.startswith('year_')]
        cohorts=[x for x in run['level_names'] if x.startswith('cohort_')]
        merged(ws,5,f"دامی‌ها: {len(years)} سال، ۳ فصل، {len(cohorts)} گروه ورود؛ بدون wave؛ بدون تحصیلات؛ تعداد کل پارامترهای تقاضا: {run['demand_parameters']}")
        merged(ws,6,'میانگین‌ها: قیمت، جمعیت‌شناسی منتخب، لگاریتم مخارج، میانگین مربع لگاریتم مخارج، میانگین CF' if spec['panel'] else 'در هیچ‌یک از مراحل میانگین خانواری به مدل مقطعی اضافه نشده است.')
        merged(ws,7,'در مدل تاریخی پنلی، میانگین تعامل‌های دوره×قیمت و دوره×مخارج و دوره×مربع مخارج نیز لحاظ شده است.' if spec['history'] and spec['panel'] else 'تابع کنترل: باقیماندهٔ پیش‌بینی لگاریتم مخارج با لگاریتم درآمد واقعی و مربع آن؛ درآمدها مثبت')
        merged(ws,8,'کشش‌ها مشتق مستقیم سهم پیش‌بینی‌شدهٔ S&Y هستند؛ میانگین‌های خانواری و CF هنگام مشتق‌گیری ثابت‌اند.')
        merged(ws,9,'ردیف = گروه تقاضا؛ ستون = قیمت گروه؛ قطر سبز = کشش خودی؛ ستون آخر = کشش مخارج غذایی')
        conv=run['convergence']
        merged(ws,10,f"همگرایی: {conv['converged']}؛ نسبت بهبود GN: {conv['scaled_gn_ratio']:.3g}؛ تکرار بیرونی: {conv['outer_iterations']}؛ tol GN=1e-8، بیرونی=1e-5")
        merged(ws,11,'توجه: جدول تاریخی «کل» میانگین وزن‌دار سه نقطهٔ دوره‌ای است؛ به معنی یک دورهٔ فرضی در نقطهٔ میانگین نیست.' if spec['history'] else 'نقطهٔ ارزیابی: میانگین غیرموزون لگاریتم قیمت، مخارج و کنترل‌ها؛ Φ و φ در میانگین شاخص پروبیت محاسبه شده‌اند.')
        merged(ws,12,'وزن سال جاری در مرحلهٔ درآمد→مخارج استفاده شده؛ پروبیت و IFGNLS مانند B1 قبلی غیرموزون‌اند. این جدول فاصلهٔ اطمینان نهایی ندارد.')
        row=matrix(ws,13,run['elasticities']['overall'],'کل نمونهٔ اجرا'+(' — میانگین نقاط سه دوره' if spec['history'] else ' — در نقطهٔ میانگین'))
        if spec['history']:
            for code in range(3): row=matrix(ws,row,run['elasticities'][f'period_{code}'],PERIODS[code]+' — ضرایب و نقطهٔ میانگین همان دوره')
            merged(ws,row,'جدول‌های بعدی نقطهٔ فرضی مشترک‌اند؛ ترکیب cohort و میانگین تعامل‌ها ممکن است در آن دوره مشاهده نشده باشد. برای نتیجهٔ اصلی از نقاط واقعی دوره استفاده کنید.'); row+=2
            for code in range(3): row=matrix(ws,row,run['elasticities'][f'common_reference_period_{code}'],PERIODS[code]+' — نقطهٔ مرجع مشترک کل نمونه؛ کنترل شرایط خانوار')
        merged(ws,row,'کنترل‌های سطحی: '+', '.join(run['level_names'])); row+=2
        if spec['panel']: merged(ws,row,'میانگین‌های شناسایی‌شده: '+', '.join(run['mundlak_names'])); ws.row_dimensions[row].height=100
        ws.column_dimensions['A'].width=30
        for j in range(2,14): ws.column_dimensions[get_column_letter(j)].width=17
        ws.column_dimensions['N'].width=18
        ws.print_options.horizontalCentered=True; ws.page_setup.orientation='landscape'
        ws.page_setup.paperSize=ws.PAPERSIZE_A3; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.print_area=f'A1:N{row+1}'
    path=OUT/'paneldemand_B1_fourteen_runs.xlsx'; wb.save(path)
    reopened=openpyxl.load_workbook(path,data_only=True)
    assert len(reopened.sheetnames)==14
    for run in results:
        ws=reopened[run['specification']['name']]
        for i in range(12):
            for j in range(12):
                assert abs(ws.cell(15+i,2+j).value-run['elasticities']['overall']['marshallian'][i][j])<1e-12
            assert abs(ws.cell(15+i,14).value-run['elasticities']['overall']['expenditure'][i])<1e-12
    return path


METHOD = (OUT/'methodology.md').read_text()


def report(results, manifest):
    penalized=[(r['specification']['name'],d['group'],GROUPS[d['group']-1]) for r in results
               for d in r['selection'] if d.get('estimation_method',{}).get('method')=='Jeffreys_prior_penalized_probit']
    lines=['# چهارده اجرای B1 توسعه‌یافته — ۱۴۰۵/۲۰۲۶\n',
           '## خلاصهٔ اجرایی\n',
           f"هر ۱۴ مدل اجرا و همگرا شدند. نمونهٔ All شامل **{manifest['sample']['All']['rows']:,}** مشاهده و Panel B شامل **{manifest['sample']['B']['rows']:,} مشاهده از {manifest['sample']['B']['households']:,} خانوار سه‌موجی** است. قیمت و سبد بین نمونه‌ها مشترک‌اند؛ دادهٔ منبع تغییر نکرده است.\n",
           'اکسل [چهارده اجرا](paneldemand_B1_fourteen_runs.xlsx) دقیقاً ۱۴ شیت دارد؛ هر شیت مشخصات، ماتریس کامل ۱۲×۱۲ مارشالی و کشش مخارج را دارد. نسخه‌های تاریخی علاوه بر خلاصهٔ کل، سه ماتریس دوره‌ای دارند. [results.json](results.json) تمام پارامترهای تقاضا، مشارکت، مرحلهٔ اول و کنترل‌ها را نگهداری می‌کند.\n',
           mdtable(['اجرا','نوع','مشاهده','خانوار','پارامتر تقاضا','outer','نسبت GN'],
           [[r['specification']['name'],model_description(r['specification']),r['sample']['rows'],r['sample']['households'],r['demand_parameters'],r['convergence']['outer_iterations'],r['convergence']['scaled_gn_ratio']] for r in results]),
           f'از ۱۶۸ معادلهٔ مشارکت، **{len(penalized)} معادله** به علت separation تأییدشده با Jeffreys برازش شده‌اند؛ سایر معادلات likelihood عادی دارند. مشاهده و کنترل برای این اصلاح حذف نشده است.\n',
           mdtable(['اجرا','گروه','نام گروه'],penalized),
           'دو نکتهٔ اقتصادی نیازمند توجه‌اند: کشش خودی لبنیات و تخم‌مرغ در cohort ۱۳۹۴ اندکی مثبت است (حدود **+۰٫۰۰۷۲**). همچنین نقطهٔ مرجع مشترکِ تاریخی برای بعضی ترکیب‌های cohort/میانگین تعامل، خارج از ترکیب‌های مشاهده‌شده است؛ مثلاً کشش مخارج مغزها در دورهٔ اول در آن نقطه **−۱٫۱۷۹** می‌شود، اما در نقطهٔ واقعی همان دوره **+۱٫۴۲۶** است. این دو رقم را به جای هم یا به‌عنوان نتیجهٔ قطعی تغییر رفتار خانوار گزارش نکنید.\n',
           '## مقایسهٔ کشش‌های اصلی کل نمونه\n',
           'ترتیب گروه‌ها در همهٔ جدول‌ها یکسان است. مقادیر کل مدل تاریخی میانگین وزن‌دار سه نقطهٔ دوره‌ای‌اند.\n',
           mdtable(['گروه']+[r['specification']['name'] for r in results],
           [[GROUPS[i]]+[r['elasticities']['overall']['marshallian'][i][i] for r in results] for i in range(12)]),
           '\nکشش‌های مخارج غذایی:\n',
           mdtable(['گروه']+[r['specification']['name'] for r in results],
           [[GROUPS[i]]+[r['elasticities']['overall']['expenditure'][i] for r in results] for i in range(12)]),
           '\n## اندازهٔ اقتصادی تفاوت‌ها\n',
           'جدول زیر تفاوت پنل تجمیعی مشترک و مقطع B مشترک را در همان نمونه و قیمت نشان می‌دهد. این تفاوت توصیفی است؛ مطابق توضیح روش، RF و کنترل سن هم تغییر می‌کنند و نسبت‌دادن همهٔ آن به موندلاک درست نیست.\n',
           mdtable(['گروه','تفاوت کشش خودی پنل−مقطع B','تفاوت کشش مخارج پنل−مقطع B'],
                   [[GROUPS[i],economic_comparisons(results)['Panel_B_vs_CS_B']['own_price_difference'][i],
                     economic_comparisons(results)['Panel_B_vs_CS_B']['expenditure_difference'][i]] for i in range(12)]),
           METHOD,
           '\n## جریان نمونه و قیمت\n']
    for name,f in manifest['sample'].items():
        lines += [f'### {name}\n',f"شروع: {f['source_rows']:,} ردیف؛ پایان: {f['rows']:,} ردیف.\n",
                  mdtable(['معیار','حذف جدید','باقی‌مانده'],[[x['criterion'],x['excluded'],x['remaining']] for x in f['flow']]),
                  f"SHA256 شناسه‌های نمونه: `{f['sample_hash']}`\n"]
    audit=json.loads((WORK/'price_audit.json').read_text())
    lines += ['قیمت ۱۰۶ قلم از All ساخته شده: Majumder کیفیت، median بازار، حداقل ۳ قیمت معتبر، سلسلهٔ county-season-UR → county-season → province-season-UR → province-season → region-season → national-season. وزن Young ثابت ملی از مخارج واقعی All و سهم مساوی سال‌هاست؛ normalization همان مرکز مشترک All است. قیمت‌ها برای buyers و non-buyers یکسان‌اند. بازسازی ماتریس Panel B با کش قبلی در tolerance مطلق 2e-12 برابر بود؛ fallback ملیِ ناموفق صفر بود.\n',
              mdtable(['Tier','All درصد','Panel B درصد'],[[str(t),audit['tiers']['All']['overall'][str(t)]*100,audit['tiers']['PanelB']['overall'][str(t)]*100] for t in range(1,7)])]
    for name in ('All','PanelB'):
        lines += [f'پوشش گروهی قیمت {name} (درصد):\n',mdtable(['گروه']+[str(t) for t in range(1,7)],[[GROUPS[i]]+[audit['tiers'][name]['by_group'][str(i+1)][str(t)]*100 for t in range(1,7)] for i in range(12)])]
    comparison=economic_comparisons(results)
    cc=comparison['separate_cohorts']
    lines += ['\n## دامنه و میانگین هشت پنل مستقل\n',
              'میانگین cohortها مدل نهم یا estimator جدید نیست؛ خلاصهٔ توصیفی هشت نقطهٔ برآورد مستقل است. وزن household count، وزن survey نیست.\n',
              mdtable(['گروه','خودی min','خودی max','خودی میانگین ساده','خودی میانگین وزن تعداد خانوار','مخارج min','مخارج max','مخارج میانگین وزن تعداد خانوار'],[[GROUPS[i],cc['own_min'][i],cc['own_max'][i],cc['own_equal_cohort_mean'][i],cc['own_household_count_weighted_mean'][i],cc['expenditure_min'][i],cc['expenditure_max'][i],cc['expenditure_household_count_weighted_mean'][i]] for i in range(12)])]
    for r in results:
        if not r['specification']['history']: continue
        v=comparison[r['specification']['name']]
        lines += [f"\n## تفاوت تاریخی در نقطهٔ مشترک: {r['specification']['name']}\n",
                  'در این جدول شرایط اقتصادی و کنترل‌ها ثابت‌اند؛ پاسخ مشارکت و ضرایب تقاضا در دوره‌ها متفاوت است. دامنه max−min یک اندازهٔ اقتصادی است، نه آزمون آماری یا یک آستانهٔ تصمیم خودکار.\n',
                  mdtable(['گروه']+['خودی '+p for p in PERIODS]+['دامنهٔ خودی']+['مخارج '+p for p in PERIODS]+['دامنهٔ مخارج'],[[GROUPS[i]]+[v['common_reference_own_by_period'][c][i] for c in range(3)]+[v['own_range'][i]]+[v['common_reference_expenditure_by_period'][c][i] for c in range(3)]+[v['expenditure_range'][i]] for i in range(12)])]
    lines += ['\n## جزئیات هر اجرا: پارامترها، کشش‌ها و diagnostics\n']
    for r in results:
        sp=r['specification']; lines += [f"### {sp['name']}\n",model_description(sp)+'\n',
            'کنترل‌های سطحی: `'+', '.join(r['level_names'])+'`.\n',
            'میانگین‌ها: `'+', '.join(r['mundlak_names'])+'`.\n',
            'ترجمهٔ alpha و بارگذاری CF جدا؛ اندازهٔ مرحلهٔ اول: '+str(r['first_stage']['columns'])+'؛ اندازهٔ هر پروبیت: '+str(len(r['selection_names'])+1)+'.\n',
            mdtable(['شاخص','مقدار'],[['R² RF',r['first_stage']['r_squared']],['partial R² درآمد جاری',r['first_stage']['excluded_current_income_partial_r_squared']],['F کلاسیک درآمد جاری',r['first_stage']['excluded_current_income_classical_f']],['objective / observation',r['convergence']['objective_per_row']],['RMSE سهم‌ها',r['convergence']['system_rmse']],['logdet Sigma',r['convergence']['sigma_log_determinant']],['Gaussian quasi log likelihood',r['convergence']['gaussian_quasi_log_likelihood']],['حداقل eigenvalue اطلاعات مقیاس‌شده',r['convergence']['scaled_information_min_eigenvalue']],['condition اطلاعات',r['convergence']['scaled_information_condition']],['جمع متوسط سهم پیش‌بینی‌شده',r['convergence']['prediction_sum_mean']]]),
            mdtable(['گروه','روش مشارکت','مشارکت','سهم پیش‌بینی منفی','سهم پیش‌بینی بالای ۱'],[[GROUPS[i],r['selection'][i].get('estimation_method',{}).get('method','native_unpenalized_probit'),r['selection'][i]['participation_rate'],r['convergence']['prediction_negative_fraction_by_group'][i],r['convergence']['prediction_above_one_fraction_by_group'][i]] for i in range(12)])]
        for key,e in r['elasticities'].items():
            lines += [f'**کشش {key}**\n',mdtable(['گروه']+GROUPS+['مخارج'],[[GROUPS[i]]+e['marshallian'][i]+[e['expenditure'][i]] for i in range(12)])]
        for code,c in r['coefficients'].items():
            lines += [f'**پارامترهای هسته، دوره {code}**\n',mdtable(['گروه','alpha مرکز','beta','lambda','delta','CF جاری'],[[GROUPS[i],c['alpha'][i],c['beta'][i],c['lambda'][i],c['delta'][i],c['current_cf'][i]] for i in range(12)]),
                      mdtable(['Gamma']+GROUPS,[[GROUPS[i]]+c['Gamma'][i] for i in range(12)])]
        lines += ['**ضرایب کنترل‌ها در واحد اصلی** (ستون mean_cf افزودهٔ داخل Phi است؛ بقیه ترجمهٔ alpha):\n',
                  mdtable(['کنترل']+GROUPS,[[n]+r['control_coefficients_original_units'][i] for i,n in enumerate(r['control_names'])]),
                  '**ضرایب مشارکت در واحد اصلی؛ ثابت آخر است:**\n',
                  mdtable(['متغیر']+GROUPS,[[n]+[r['selection_coefficients'][i][j] for i in range(12)] for j,n in enumerate(r['selection_names']+['constant'])]),
                  '**تمام ضرایب RF:**\n',mdtable(['متغیر','ضریب'],list(r['first_stage']['coefficients'].items())),
                  '**حذف ستون‌های همخط/ثابت:**\n', '```json\n'+json.dumps(r['removed_columns'],ensure_ascii=False,indent=2)+'\n```\n']
        if 'within_between' in r:
            lines += ['**تغییرات درون/بین خانوار:**\n',mdtable(['متغیر','SD کل','SD بین','SD درون','سهم واریانس درون'],[[k,v['overall_sd'],v['between_sd'],v['within_sd'],v['within_variance_fraction']] for k,v in r['within_between'].items()]),
                      '**تغییر وضعیت خرید:**\n',mdtable(['گروه','همیشه صفر','همیشه مثبت','تغییر وضعیت'],[[GROUPS[i],r['switching'][str(i+1)]['always_zero'],r['switching'][str(i+1)]['always_positive'],r['switching'][str(i+1)]['switchers']] for i in range(12)])]
            wd=r['within_prices_after_time_controls']
            lines += ['**variation قیمت درون خانوار پس از جذب سال/فصل:**\n',
                      mdtable(['قیمت','SD باقیماندهٔ درون','سهم variance درونِ باقی‌مانده'],[[k,v['within_sd_after_time_effects'],v['fraction_of_within_variance_remaining']] for k,v in wd['prices'].items()]),
                      f"رتبهٔ قیمت‌های باقیمانده: {wd['residual_price_rank']}؛ condition: {wd['residual_price_condition']:.6g}. عبور خانوار از مرز تاریخی: `{r['households_crossing_boundaries']}`.\n"]
        if 'conditional_historical_tests' in r:
            lines += ['**آزمون‌های تاریخی مشروط، غیرنهایی:**\n',mdtable(['بلوک','Wald','df','p'],[[k,v['statistic'],v['df'],v['pvalue']] for k,v in r['conditional_historical_tests'].items()])]
    lines += ['\n## بازتولید\n','```bash\npython -m pip install -r pilot/requirements.txt\npython -m pilot.b1_fourteen.run --prepare\npython -m unittest discover -s pilot/tests -p "test_*.py"\npython -m pilot.b1_fourteen.run --run-all --workers 2\npython -m pilot.b1_fourteen.selection_audit\npython -m pilot.b1_fourteen.run --run-all --workers 2\npython -m pilot.b1_fourteen.run --refresh-results\npython -m pilot.b1_fourteen.conditional_tests CS_All_history\npython -m pilot.b1_fourteen.conditional_tests CS_B_history\npython -m pilot.b1_fourteen.conditional_tests Panel_B_history\npython -m pilot.b1_fourteen.export\n```\n',
              'نسخهٔ مدل/manifest: `'+VERSION+'`. داده‌ها و fitted objects در intermediate باقی‌اند؛ در Git فقط این بستهٔ کد، manifest، گزارش، نتایج تجمیعی و اکسل قرار دارد.\n']
    (OUT/'report.md').write_text('\n'.join(lines))


def main():
    results=[json.loads((WORK/f"{m['name']}_result.json").read_text()) for m in MODELS]
    for r in results:
        assert r['version']==VERSION and r['convergence']['converged']
        assert r['rank_selection']['rank']==r['rank_selection']['columns']
        for e in r['elasticities'].values():
            assert np.asarray(e['marshallian']).shape==(12,12)
            assert np.isfinite(e['marshallian']).all() and np.isfinite(e['expenditure']).all()
            if 'max_derivative_absolute_error' in e: assert e['max_derivative_absolute_error']<1e-6
    manifest=json.loads((OUT/'manifest.json').read_text())
    audit=json.loads((WORK/'price_audit.json').read_text())
    dump(OUT/'results.json',{'version':VERSION,'manifest':manifest,'groups':GROUPS,
         'price_tiers':audit['tiers'],'economic_comparisons':economic_comparisons(results),'runs':results})
    path=workbook(results); report(results,manifest)
    print('EXPORTED',len(results),'runs',path)


if __name__=='__main__': main()
