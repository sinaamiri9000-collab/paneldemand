"""Extend the supplied own-price workbook with complete Marshallian matrices.

Original sheets, summary values and styles are retained. New numbers come from
the unrounded fixed-parameter results, never from rounded Markdown tables.
Formula caches allow previews to display ranges before Excel recalculates them.
"""
from copy import copy
from pathlib import Path
import json
import zipfile
import xml.etree.ElementTree as ET

import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

OUT = Path(__file__).resolve().parent
TEMPLATE = OUT / 'own_price_template.xlsx'
DEST = OUT / 'paneldemand_B1_B2_marshallian_12x12_1392_1403.xlsx'
MODELS = ('B1_no_season_means', 'B2_no_season_means')
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def patch_caches(path, sheetnames, caches):
    """Write numeric formula results without changing the formulas or styling."""
    with zipfile.ZipFile(path) as z:
        files = [(info, z.read(info.filename)) for info in z.infolist()]
    for index, name in enumerate(sheetnames, 1):
        filename = f'xl/worksheets/sheet{index}.xml'
        position = next(i for i, (info, _) in enumerate(files) if info.filename == filename)
        info, data = files[position]
        root = ET.fromstring(data)
        for cell in root.findall('.//s:c', NS):
            if cell.find('s:f', NS) is None:
                continue
            key = (name, cell.attrib['r'])
            value = caches[key]
            assert value is not None and np.isfinite(value), key
            node = cell.find('s:v', NS)
            if node is None:
                node = ET.SubElement(cell, '{'+NS['s']+'}v')
            node.text = repr(float(value))
        files[position] = (info, ET.tostring(root, encoding='utf-8', xml_declaration=True))
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        for info, data in files:
            z.writestr(info, data)


def build():
    data = json.loads((OUT / 'results.json').read_text())
    groups = data['groups']
    book = load_workbook(TEMPLATE)
    template_values = load_workbook(TEMPLATE, data_only=True)
    base = book['کل نمونه']
    caches = {(s.title, c.coordinate): template_values[s.title][c.coordinate].value
              for s in book for row in s for c in row if c.data_type == 'f'}
    originals = {(s.title, c.coordinate): c.value for s in book for row in s
                 for c in row if c.value is not None}
    locations = {}

    def banner(sheet, row, text, b2=False):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=13)
        c = sheet.cell(row, 1, text)
        c._style = copy(base['D4' if b2 else 'A4']._style)
        c.alignment = Alignment(vertical='center', horizontal='left', readingOrder=2)
        sheet.row_dimensions[row].height = 28

    def matrix(sheet, start, title, numbers, model_index, formulas=None):
        banner(sheet, start, title, model_index == 1)
        sheet.cell(start+1, 1, 'گروه تقاضا ↓ / گروه قیمت →')
        for j, group in enumerate(groups, 2):
            sheet.cell(start+1, j, group)
        for cell in sheet[start+1][:13]:
            cell._style = copy(base['A5']._style)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True, readingOrder=2)
        sheet.row_dimensions[start+1].height = 58
        for i, group in enumerate(groups):
            row = start+2+i
            label = sheet.cell(row, 1, group)
            label._style = copy(base.cell(6+i, 1)._style)
            label.alignment = Alignment(vertical='center', wrap_text=True, readingOrder=2)
            sheet.row_dimensions[row].height = 30
            for j in range(12):
                value = float(numbers[i, j])
                c = sheet.cell(row, j+2, value if formulas is None else formulas[i][j])
                c._style = copy(base.cell(6+i, 2+model_index)._style)
                c.alignment = Alignment(horizontal='center', vertical='center')
                if i == j:
                    font = copy(c.font); font.bold = True; c.font = font
                if formulas is not None:
                    caches[(sheet.title, c.coordinate)] = value
        return start+2

    points = [('کل نمونه', 'global', None)]
    points += [(f'سال {key}', 'years', key) for key in data['models'][MODELS[0]]['years']]
    points += [(f'کوهرت {key}', 'cohorts', key) for key in data['models'][MODELS[0]]['cohorts']]
    assert book.sheetnames == [x[0] for x in points]+['دامنه تغییرات']
    for title, section, key in points:
        sheet = book[title]
        for k, name in enumerate(MODELS):
            model = data['models'][name]
            record = model[section] if key is None else model[section][key]
            numbers = np.asarray(record['elasticities']['marshallian'])
            # Uploaded own-price/expenditure summaries must match the same fits.
            for i in range(12):
                assert abs(sheet.cell(6+i, 2+k).value-round(numbers[i, i], 4)) < 1e-12
                assert abs(sheet.cell(6+i, 4+k).value-round(record['elasticities']['expenditure'][i], 4)) < 1e-12
            start = 24 if k == 0 else 40
            locations[(title, k)] = matrix(sheet, start, f'ماتریس کامل کشش مارشالی ۱۲×۱۲ | B{k+1}', numbers, k)
        sheet.merge_cells('A22:M22')
        sheet['A22'] = 'ردیف: تقاضای گروه • ستون: قیمت گروه • قطر پررنگ: کشش خودی • سایر خانه‌ها: کشش متقاطع'
        sheet['A22']._style = copy(base['A19']._style)
        sheet.row_dimensions[22].height = 24
        sheet.merge_cells('A55:M55')
        sheet['A55'] = 'اعداد ماتریس با دقت کامل ذخیره و با چهار رقم اعشار نمایش داده می‌شوند؛ ضرایب و نقاط مرجع عین نتایج قبلی‌اند.'
        sheet['A55']._style = copy(base['A19']._style)
        sheet.row_dimensions[55].height = 24
        for j in range(6, 14):
            sheet.column_dimensions[get_column_letter(j)].width = 18
        sheet.freeze_panes = 'B6'
        sheet.print_area = 'A1:M55'
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    sheet = book['دامنه تغییرات']
    sheet.merge_cells('A78:M78')
    sheet['A78'] = 'دامنهٔ کامل ۱۲×۱۲: کمینه، بیشینه و بیشینه منهای کمینهٔ هر خانه، با استفاده از ماتریس‌های دقیق شیت‌ها'
    sheet['A78']._style = copy(base['A19']._style)
    sheet.row_dimensions[78].height = 28
    start = 80
    range_checks = []
    for section, label in [('years', 'سالانه'), ('cohorts', 'کوهرت')]:
        titles = [p[0] for p in points if p[1] == section]
        for k, name in enumerate(MODELS):
            values = np.array([r['elasticities']['marshallian'] for r in data['models'][name][section].values()])
            for operation, caption, numbers in [('MIN', 'کمینه', values.min(0)),
                                               ('MAX', 'بیشینه', values.max(0)),
                                               ('RANGE', 'دامنه', np.ptp(values, axis=0))]:
                formulas = []
                for i in range(12):
                    row = []
                    for j in range(12):
                        refs = ','.join(f"'{t}'!{get_column_letter(j+2)}{locations[(t,k)]+i}" for t in titles)
                        row.append(f'=MAX({refs})-MIN({refs})' if operation == 'RANGE' else f'={operation}({refs})')
                    formulas.append(row)
                top = matrix(sheet, start, f'{label} • {caption} کشش مارشالی ۱۲×۱۲ | B{k+1}', numbers, k, formulas)
                range_checks.append((top, numbers))
                start += 17
    for j in range(2, 14):
        sheet.column_dimensions[get_column_letter(j)].width = 18
    sheet.column_dimensions['E'].width = 18
    sheet.print_area = f'A1:M{start-3}'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')
    book.save(DEST)
    patch_caches(DEST, book.sheetnames, caches)

    # Validate the delivered file itself, including formula caches and orientation.
    actual = load_workbook(DEST, data_only=True)
    formulas = load_workbook(DEST)
    assert actual.sheetnames == book.sheetnames
    for (title, coordinate), value in originals.items():
        assert formulas[title][coordinate].value == value, (title, coordinate)
    for title, section, key in points:
        for k, name in enumerate(MODELS):
            record = data['models'][name][section] if key is None else data['models'][name][section][key]
            top = locations[(title, k)]
            matrix_values = [[actual[title].cell(top+i, j+2).value for j in range(12)] for i in range(12)]
            np.testing.assert_allclose(matrix_values, record['elasticities']['marshallian'], rtol=0, atol=1e-14)
    for top, numbers in range_checks:
        values = [[actual['دامنه تغییرات'].cell(top+i, j+2).value for j in range(12)] for i in range(12)]
        np.testing.assert_allclose(values, numbers, rtol=0, atol=1e-14)
    for (title, coord), value in caches.items():
        assert actual[title][coord].value == value, (title, coord)
    print(f'Validated 22 sheets, 42 complete matrices, 12 range matrices, all original cells and {len(caches)} cached formulas: {DEST}')


if __name__ == '__main__':
    build()
