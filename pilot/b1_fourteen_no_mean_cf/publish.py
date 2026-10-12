"""Release ten no-mean-CF panels and four unchanged CS records in one workbook."""
from .run import base, OUT, WORK, VERSION, NAMES, PREVIOUS, REMOVED, prepare, is_square_mean
from pilot.b1_fourteen.export import workbook, mdtable
from pilot.cohort1401_nuts_sensitivity.run import clean, save_json
import json
import shutil
import numpy as np
import openpyxl


def report(runs, comparisons):
    lines = ['# بازتخمین ده پنل با حذف میانگین CF\n',
        'ده مدل پنلی نسخهٔ بدون میانگین مربع مخارج دوباره تخمین زده شدند. فقط `mean_cf_residual` از هر دوازده معادلهٔ مشارکت و از کنترل‌های تقاضای نهفته حذف شده است. CF جاری، میانگین لگاریتم مخارج، سایر میانگین‌ها، دامی‌ها، فرم QUAIDS/S&Y، قیود و روش محاسبهٔ کشش حفظ شده‌اند. چهار مدل مقطعی به علت نداشتن میانگین CF عیناً بازاستفاده شده‌اند.\n',
        'مرحلهٔ اول مخارج هیچ میانگین CF ندارد و تغییر نمی‌کند: همان RF وزن‌دار و همان باقیماندهٔ جاری از نسخهٔ قبل استفاده شده است. مشارکت و IFGNLS مانند مدل A قبلی بدون وزن‌اند. میانگین خانوار بدون وزن سه موج است. حذف میانگین CF به معنی حذف تابع کنترل جاری نیست.\n',
        'اصلاح مشارکت نان cohortهای ۱۳۹۸ و ۱۴۰۱ و سبزیجات cohort ۱۳۹۹، یعنی حذف Bachelor و میانگین آن فقط از همان معادله، حفظ شده است. دادهٔ منبع و قیمت‌ها تغییر نکرده‌اند؛ هیچ bootstrap اجرا نشده است.\n',
        '[نتایج کامل و همهٔ ضرایب](results.json) · [اکسل فعلی ۱۴ شیت](../b1_fourteen/paneldemand_B1_fourteen_runs.xlsx) · [نتایج مبنای مقایسه](../b1_fourteen_no_exp_square/results.json). [اکسل قبلی با میانگین CF](https://github.com/sinaamiri9000-collab/paneldemand/blob/be4e5f48417670c5fab24f40567e58c4c941e309/pilot/b1_fourteen/paneldemand_B1_fourteen_runs.xlsx) در تاریخچه حفظ شده است.\n',
        '## همگرایی\n', mdtable(['مدل', 'مشاهده', 'خانوار', 'پارامتر تقاضا', 'تکرار بیرونی', 'نسبت GN', 'زمان تخمین تقاضا، ثانیه'],
            [[r['specification']['name'], r['sample']['rows'], r['sample']['households'], r['demand_parameters'],
              r['convergence']['outer_iterations'], r['convergence']['scaled_gn_ratio'], r['convergence'].get('point_seconds', r.get('total_seconds'))]
             for r in runs if r['specification']['panel']]),
        '\nتنظیمات حل‌کننده همان نسخهٔ قبل است: `gn_tol=1e-8` و معیار بیرونی پیش‌فرض `1e-5`. رتبهٔ طراحی، محدودیت‌های تقاضا، مثبت‌بودن اطلاعات مقیاس‌شده و نسبت GN کنترل می‌شوند؛ صرف موفقیت optimizer کافی نیست.\n',
        '## مقایسهٔ کشش‌های قبل و بعد\n']
    for name, c in comparisons.items():
        lines += [f'### {name}\n', mdtable(['گروه', 'مخارج قبل', 'مخارج جدید', 'تفاوت', 'خودی قبل', 'خودی جدید', 'تفاوت'],
            [[base.GROUPS[i], c['before_expenditure'][i], c['after_expenditure'][i], c['expenditure_difference'][i],
              c['before_own'][i], c['after_own'][i], c['own_price_difference'][i]] for i in range(12)])]
    lines += ['\n## تفسیر و بازتولید\n',
        'این مقایسه توصیفی است و آزمون معناداری نهایی تفاوت‌ها نیست. کشش‌ها از مشتق مستقیم سهم پیش‌بینی‌شدهٔ S&Y به دست می‌آیند؛ میانگین‌های خانوار و CF هنگام مشتق جاری ثابت‌اند. در مدل تاریخی، خلاصهٔ کل میانگین سه نقطهٔ دوره‌ای با وزن تعداد مشاهده است؛ ماتریس‌های دوره‌ای و نقاط مرجع مشترک نیز در اکسل هستند. سهم پیش‌بینی‌شدهٔ غیرمثبت، کشش نامعتبر تولید می‌کند و عدد جایگزین نمی‌گیرد.\n',
        '```bash\npython -m pilot.b1_fourteen_no_mean_cf.run --run-all --workers 2\npython -m pilot.b1_fourteen_no_mean_cf.publish\npython -m unittest pilot.b1_fourteen_no_mean_cf.test_checks\n```\n',
        'دو subprocess مستقل، هر کدام با یک BLAS thread، استفاده شده‌اند. اشیای برازش، داده و logها در `intermediate/pilot_inputs/b1_fourteen_no_mean_cf/` خصوصی باقی می‌مانند. کد، گزارش، نتایج و اکسل در مخزن ثبت شده‌اند.\n']
    (OUT / 'report.md').write_text('\n'.join(lines))


def publish():
    manifest = prepare()
    previous = json.loads(PREVIOUS.read_text())
    before = {r['specification']['name']: r for r in previous['runs']}
    runs = []; comparisons = {}
    for spec in base.MODELS:
        name = spec['name']; old = before[name]
        if not spec['panel']:
            runs.append(old); continue
        r = json.loads((WORK / f'{name}_result.json').read_text())
        assert r['version'] == VERSION and r['convergence']['converged']
        assert r['sample'] == old['sample']
        assert r['demand_parameters'] == old['demand_parameters']-11
        assert r['control_names'] == [n for n in old['control_names'] if n != REMOVED]
        assert REMOVED not in r['selection_names']
        assert not any(is_square_mean(n) for n in r['control_names'] + r['selection_names'])
        assert r['first_stage'] == old['first_stage']
        assert [d.get('excluded_names', []) for d in r['selection']] == [d.get('excluded_names', []) for d in old['selection']]
        a = r['elasticities']['overall']; b = old['elasticities']['overall']
        comparisons[name] = clean(dict(before_expenditure=b['expenditure'], after_expenditure=a['expenditure'],
            expenditure_difference=np.asarray(a['expenditure'], float)-np.asarray(b['expenditure'], float),
            before_own=np.diag(np.asarray(b['marshallian'], float)), after_own=np.diag(np.asarray(a['marshallian'], float)),
            own_price_difference=np.diag(np.asarray(a['marshallian'], float)-np.asarray(b['marshallian'], float)),
            marshallian_difference=np.asarray(a['marshallian'], float)-np.asarray(b['marshallian'], float)))
        runs.append(r)
    assert len(runs) == 14 and runs[:4] == previous['runs'][:4]
    save_json(OUT / 'results.json', dict(version=VERSION, groups=base.GROUPS, manifest=manifest,
        runs=runs, comparisons=comparisons, refitted_models=NAMES,
        unchanged_cross_sections=[r['specification']['name'] for r in runs[:4]],
        bootstrap_run=False, QUAIDS_and_SY_unchanged=True,
        note='Historical excluded-income diagnostic erratum remains applicable to the four unchanged CS records; see pilot/README.md.'))
    oldbook = openpyxl.load_workbook(base.OUT / 'paneldemand_B1_fourteen_runs.xlsx', data_only=True)
    temp = WORK / 'paneldemand_B1_fourteen_runs.xlsx'
    workbook(runs, temp)
    newbook = openpyxl.load_workbook(temp, data_only=True)
    for r in runs[:4]:
        name = r['specification']['name']
        assert list(oldbook[name].values) == list(newbook[name].values), name
    for r in runs[4:]:
        ws = newbook[r['specification']['name']]
        assert 'میانگین CF حذف شده' in ws['A2'].value
        assert 'میانگین مربع لگاریتم مخارج حذف شده' in ws['A2'].value
        assert not any(REMOVED in str(c.value) for row in ws for c in row)
    for ws in newbook:
        assert ws.freeze_panes is None and ws.sheet_view.pane is None
        assert ws.sheet_view.zoomScale == 85
        assert max(d.height or 0 for d in ws.row_dimensions.values()) <= 36
    report(runs, comparisons)
    shutil.copy2(temp, base.OUT / 'paneldemand_B1_fourteen_runs.xlsx')
    print('UPDATED', base.OUT / 'paneldemand_B1_fourteen_runs.xlsx', flush=True)


if __name__ == '__main__': publish()
