"""Verify only the additive mean-CF effect is removed, including historical fits."""
import copy
import json
import tempfile
import unittest
from pathlib import Path
import numpy as np
from scipy.stats import norm
from pyquaidsce.model import DemandData
from pyquaidsce.params import unpack
from .run import base, trim_controls, warm_start, REMOVED
from pilot.b1_fourteen.export import workbook


class Checks(unittest.TestCase):
    def test_mean_cf_removal_preserves_all_other_effects(self):
        rng = np.random.default_rng(1541)
        for history in (False, True):
            with self.subTest(history=history):
                n = 36; m = 12
                names = ['household_size', 'mean_ln_expenditure', REMOVED]
                Z = rng.normal(0, .15, (n, 3)); centers = rng.normal(size=3)
                centers[-1] = 0; scales = np.exp(rng.normal(size=3))
                lp = rng.normal(0, .1, (n, m)); y = rng.normal(3, .2, n)
                cf = rng.normal(0, .1, n)
                sn = base.PN + ['ln_expenditure'] + names + ['cf_residual']
                X = np.column_stack([lp, y, Z*scales+centers, cf])
                tau = rng.normal(0, .02, (m, len(sn)+1)); k = X@tau[:, :-1].T+tau[:, -1]
                data = DemandData(lp, y, rng.dirichlet(np.ones(m)*3, n), np.zeros((n, 1)),
                    norm.cdf(k), norm.pdf(k), 1., cf)
                old = dict(data=data, Z=Z, centers=centers, scales=scales, control_names=names,
                    selection_names=sn, tau=tau, selection_index=k, regimes=np.repeat(np.arange(3), 12),
                    additive_columns=[2], mundlak_names=names[1:], rf_names=['constant', 'ln_income_centered'],
                    Xrf=rng.normal(size=(n, 2)), specification={'history': history})
                oldcore = base.core_for(old)
                theta = rng.normal(0, .002, oldcore.spec.n_free)
                theta[oldcore.base_slices['alpha']] = 1/m
                fit = type('Fit', (), {'theta': theta})()
                new = copy.deepcopy(old); reduced_X = trim_controls(new)
                newcore = base.core_for(new); th = warm_start(newcore, new, fit, old)
                self.assertEqual(oldcore.spec.n_free-newcore.spec.n_free, 11)
                np.testing.assert_array_equal(new['Xrf'], old['Xrf'])
                np.testing.assert_array_equal(new['data'].control_function, cf)
                np.testing.assert_array_equal(reduced_X, X[:, [j for j, name in enumerate(sn) if name != REMOVED]])
                for code in range(3 if history else 1):
                    ot, oe = oldcore.unpack(theta, code); nt, ne = newcore.unpack(th, code)
                    op = unpack(ot, oldcore.native); np_ = unpack(nt, newcore.native)
                    for field in ('alpha', 'beta', 'gamma', 'lam', 'delta', 'cfcoef'):
                        np.testing.assert_allclose(getattr(op, field), getattr(np_, field), atol=2e-16)
                    np.testing.assert_allclose(ne, oe[:2], rtol=0, atol=1e-18)
                removed_effect = data.cdf * Z[:, 2, None] * oldcore.unpack(theta)[1][2]
                np.testing.assert_allclose(newcore.fitted(th), oldcore.fitted(theta)-removed_effect, atol=2e-16)

    def test_workbook_labels_mean_cf_from_actual_controls(self):
        results = json.loads((base.OUT / 'results.json').read_text())['runs']
        for r in results:
            if r['specification']['panel']:
                for key in ('control_names', 'mundlak_names'):
                    r[key] = [n for n in r[key] if n != REMOVED and not n.endswith('centered_ln_expenditure_squared')]
        with tempfile.TemporaryDirectory() as td:
            path = workbook(results, Path(td) / 'test.xlsx')
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True)
            for r in results:
                ws = wb[r['specification']['name']]
                if r['specification']['panel']:
                    self.assertIn('میانگین CF حذف شده', ws['A2'].value)
                    self.assertFalse(any(REMOVED in str(c.value) for row in ws for c in row))
                else:
                    self.assertNotIn('میانگین CF حذف شده', ws['A2'].value)
                self.assertIsNone(ws.freeze_panes)


if __name__ == '__main__': unittest.main()
