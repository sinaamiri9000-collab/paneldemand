"""Merge ten new panels with four unchanged CS fits and update the 14-sheet XLSX."""
from .run import base,OUT,WORK,VERSION,NAMES,MODELS,prepare,sha,is_square_mean
from pilot.cohort1401_nuts_sensitivity.run import clean,save_json
from pilot.b1_fourteen.export import workbook,mdtable,model_description
import json
import shutil
from pathlib import Path
import numpy as np
import openpyxl


def report(runs,comparisons):
    lines=['# حذف میانگین مربع لگاریتم مخارج از ۱۰ مدل پنلی\n',
        '۱۰ مدل پنلی بازتخمین شده‌اند؛ چهار مدل مقطعی عین نتایج قبلی باقی مانده‌اند. میانگین خودِ لگاریتم مخارج، CF جاری و میانگین آن، سایر میانگین‌ها و دامی‌ها حفظ شده‌اند. فرم درجه‌دوم QUAIDS، S&Y، قیود هسته، نمونه، قیمت و تنظیمات حل‌کننده تغییر نکرده‌اند.\n',
        'از پنل تاریخی، کنترل `mean_centered_ln_expenditure_squared` و دو `mean_history_*_centered_ln_expenditure_squared`، در مجموع سه کنترل، حذف شده‌اند. این حذف شامل lambda یا جملهٔ درجه‌دوم مخارج جاری در QUAIDS نیست. RF و باقیمانده‌های آن ثابت مانده‌اند؛ شاخص‌های درآمد RF دوباره محاسبه شده‌اند.\n',
        'اصلاح قبلی Bachelor/mean Bachelor فقط در مشارکت نان cohort ۱۳۹۸، سبزیجات cohort ۱۳۹۹ و نان cohort ۱۴۰۱ حفظ شده است. مقایسهٔ اثر حذف مربع با همان نسخهٔ اصلاح‌شده انجام می‌شود؛ ارقام اکسل اولیه نیز جداگانه ذخیره شده‌اند.\n',
        '[اکسل به‌روز ۱۴ شیت](../b1_fourteen/paneldemand_B1_fourteen_runs.xlsx) شامل تمام ماتریس‌های ۱۲×۱۲ مارشالی و کشش‌های مخارج است. مشخصات هر شیت حذف کنترل مربع را صریحاً ثبت می‌کند. [results.json](results.json) تمام ضرایب، کنترل‌ها، ماتریس‌ها و تفاوت‌های قبل/بعد را دارد.\n',
        'نسخهٔ قبل از این حذف: [اکسل تاریخی](https://github.com/sinaamiri9000-collab/paneldemand/blob/dfc9abd6c84e4326a90c1faf2490826dfc6b1796/pilot/b1_fourteen/paneldemand_B1_fourteen_runs.xlsx). گزارش و JSON بستهٔ اولیه حفظ شده‌اند.\n',
        '## همگرایی و نمونه\n',mdtable(['مدل','مشاهده','خانوار','پارامتر','outer','نسبت GN'],
        [[r['specification']['name'],r['sample']['rows'],r['sample']['households'],r['demand_parameters'],
          r['convergence']['outer_iterations'],r['convergence']['scaled_gn_ratio']] for r in runs if r['specification']['panel']]),
        '\n## کشش مخارج مغزها، خشکبار و خرما\n',mdtable(['مدل','اکسل اولیه','قبل با اصلاح مشارکت','جدید','تفاوت حذف مربع'],
        [[n,c['initial_expenditure'][8],c['before_expenditure'][8],c['after_expenditure'][8],
          c['expenditure_difference'][8]] for n,c in comparisons.items()]),
        '\n## مقایسهٔ تمام گروه‌ها در هر مدل\n']
    for r in runs:
        if not r['specification']['panel']: continue
        name=r['specification']['name']; c=comparisons[name]
        lines += [f'### {name}\n',model_description(r['specification'])+'\n',
            'کنترل‌های حذف‌شده: `'+', '.join(r['removed_squared_expenditure_means'])+'`.\n',
            mdtable(['گروه','مخارج قبل','مخارج جدید','تفاوت','خودی قبل','خودی جدید','تفاوت'],
                [[base.GROUPS[i],c['before_expenditure'][i],c['after_expenditure'][i],c['expenditure_difference'][i],
                  c['before_own'][i],c['after_own'][i],c['own_price_difference'][i]] for i in range(12)])]
    lines+=['\n## تفسیر و بازتولید\n',
        'کشش‌ها مشتق مستقیم سهم پیش‌بینی‌شدهٔ S&Y هستند. میانگین‌های خانوار و CF در مشتق جاری ثابت‌اند. در مدل تاریخی، جدول کل میانگین با وزن تعداد مشاهدهٔ سه نقطهٔ دوره‌ای است؛ جدول‌های فرضیِ نقطهٔ مشترک همچنان ممکن است خارج از ترکیب‌های مشاهده‌شده باشند. کشش با سهم پیش‌بینی‌شدهٔ غیرمثبت «نامعتبر» نمایش داده می‌شود و عددی جایگزین آن نشده است.\n',
        'وزن سال جاری فقط در RF است؛ مشارکت و IFGNLS مانند مدل قبلی غیرموزون‌اند. هیچ bootstrap یا آزمون نهاییِ معناداری اختلاف کشش اجرا نشده است.\n',
        '```bash\npython -m pilot.b1_fourteen_no_exp_square.run --run-all --workers 2\npython -m pilot.b1_fourteen_no_exp_square.publish\npython -m unittest pilot.b1_fourteen_no_exp_square.test_checks\n```\n',
        'کار در دو subprocess مستقل و هر کدام با یک BLAS thread انجام شده است. داده و stage/fit/logهای کاری در `intermediate/pilot_inputs/b1_fourteen_no_exp_square/` هستند و منتشر نمی‌شوند. checksum منابع قبلی و داده‌های منجمد بررسی می‌شود.\n']
    (OUT/'report.md').write_text('\n'.join(lines))


def publish():
    manifest=prepare()
    for p,h in manifest['source_hashes'].items(): assert sha(base.REPO/p)==h,p
    initial=json.loads((base.OUT/'results.json').read_text())
    old={r['specification']['name']:r for r in initial['runs']}
    runs=[]; comparisons={}
    for spec in base.MODELS:
        name=spec['name']
        if not spec['panel']:
            runs.append(old[name]); continue
        run=json.loads((WORK/f'{name}_result.json').read_text())
        assert run['version']==VERSION and run['convergence']['converged']
        assert not any(is_square_mean(n) for n in run['control_names']+run['selection_names'])
        assert 'mean_ln_expenditure' in run['control_names']
        assert old[name]['sample']==run['sample']
        assert run['demand_parameters']==old[name]['demand_parameters']-(33 if spec['history'] else 11)
        before=old[name]
        corrected=base.INPUT/'marital_participation_sensitivity'/f'{name}_result.json'
        if corrected.exists(): before=json.loads(corrected.read_text())
        b=before['elasticities']['overall']; a=run['elasticities']['overall']
        orig=old[name]['elasticities']['overall']
        comparisons[name]=clean({
            'initial_expenditure':orig['expenditure'],'initial_own':np.diag(np.asarray(orig['marshallian'],float)),
            'before_expenditure':b['expenditure'],'after_expenditure':a['expenditure'],
            'expenditure_difference':np.asarray(a['expenditure'],float)-np.asarray(b['expenditure'],float),
            'before_own':np.diag(np.asarray(b['marshallian'],float)),
            'after_own':np.diag(np.asarray(a['marshallian'],float)),
            'own_price_difference':np.diag(np.asarray(a['marshallian'],float)-np.asarray(b['marshallian'],float)),
            'marshallian_difference':np.asarray(a['marshallian'],float)-np.asarray(b['marshallian'],float),
            'comparison_baseline':'previous participation-corrected model' if corrected.exists() else 'original fourteen-model fit'})
        runs.append(run)
    assert len(runs)==14
    for r in runs[:4]: assert r==old[r['specification']['name']]
    save_json(OUT/'results.json',{'version':VERSION,'groups':base.GROUPS,'manifest':manifest,
        'runs':runs,'comparisons':comparisons,'refitted_models':NAMES,
        'unchanged_cross_sections':[r['specification']['name'] for r in runs[:4]],
        'bootstrap_run':False,'QUAIDS_and_SY_unchanged':True,
        'note':'four original CS records unchanged, including historical RF diagnostic fields; their old excluded-income diagnostics carry the erratum in pilot/README.md'})
    released=json.loads((OUT/'results.json').read_text())
    assert released['runs'][:4]==initial['runs'][:4]
    # Validate all four original CS worksheets before replacing the workbook.
    oldbook=openpyxl.load_workbook(base.OUT/'paneldemand_B1_fourteen_runs.xlsx',data_only=True)
    temp=WORK/'paneldemand_B1_fourteen_runs.xlsx'
    workbook(runs,temp)
    newbook=openpyxl.load_workbook(temp,data_only=True)
    for r in runs[:4]:
        n=r['specification']['name']
        assert list(oldbook[n].values)==list(newbook[n].values),n
    for ws in newbook:
        assert ws.freeze_panes is None and ws.sheet_view.pane is None
        assert ws.sheet_view.zoomScale==85
        assert max(d.height or 0 for d in ws.row_dimensions.values())<=36
        assert max(c.font.sz or 0 for row in ws for c in row if c.value is not None)<=10
    for r in runs[4:]:
        ws=newbook[r['specification']['name']]
        assert 'میانگین مربع لگاریتم مخارج حذف شده' in ws['A2'].value
        assert not any('mean_centered_ln_expenditure_squared' in str(c.value) or
            'mean_history_1_centered_ln_expenditure_squared' in str(c.value) or
            'mean_history_2_centered_ln_expenditure_squared' in str(c.value) for row in ws for c in row)
    shutil.copy2(temp,base.OUT/'paneldemand_B1_fourteen_runs.xlsx')
    delivery=base.OUT/'paneldemand_B1_fourteen_runs.xlsx'
    report(runs,comparisons)
    print('UPDATED 14-SHEET WORKBOOK',delivery,flush=True)


if __name__=='__main__': publish()
